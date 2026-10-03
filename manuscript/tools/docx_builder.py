"""Build the manuscript .docx with the formatting of the TRI-GraphRAG manuscript (v11).

Purpose: reuse the styles, numbering, footers and page setup of the reference manuscript as a template and write new
content with the same paragraph, caption and table conventions.
Author: Jose D. Hernandez-Betancur
Date: 2026-10-03
"""

import logging
import re
from pathlib import Path
from typing import Any, Final

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

logger = logging.getLogger(__name__)

TEMPLATE: Final = Path(
    "/Users/josehernandez/Documents/tri_chem_graphrag/manuscript/TRI_GraphRAG_Manuscript_v11_JH_response_to_GRM.docx"
)
MARKUP: Final = re.compile(r"\{(sub|sup|i|b|hl):([^{}]*)\}")
DOUBLE_SPACE: Final = 480
HEADING_NUM_ID: Final = "1"


def open_template(path: Path = TEMPLATE) -> Any:
    """Open the reference manuscript and strip its body, keeping styles, numbering, footers and page setup.

    Args:
        path: Reference .docx.

    Returns:
        An empty python-docx document carrying the reference formatting.
    """
    document = Document(str(path))
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    settings = document.settings.element
    for tag in ("w:trackRevisions",):
        for element in settings.findall(qn(tag)):
            settings.remove(element)
    return document


def add_runs(paragraph: Any, text: str, size_pt: float | None = None, bold: bool = False) -> None:
    """Append text with light inline markup: {sub:..}, {sup:..}, {i:..}, {b:..}, {hl:..}.

    Args:
        paragraph: Target paragraph.
        text: Text with markup.
        size_pt: Optional font size.
        bold: Make every run bold.
    """
    position = 0
    for match in MARKUP.finditer(text):
        _plain_run(paragraph, text[position : match.start()], size_pt, bold)
        kind, content = match.groups()
        run = paragraph.add_run(content)
        _style_run(run, size_pt, bold)
        if kind == "sub":
            run.font.subscript = True
        elif kind == "sup":
            run.font.superscript = True
        elif kind == "i":
            run.italic = True
        elif kind == "b":
            run.bold = True
        elif kind == "hl":
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        position = match.end()
    _plain_run(paragraph, text[position:], size_pt, bold)


def _style_run(run: Any, size_pt: float | None, bold: bool) -> None:
    if size_pt:
        run.font.size = Pt(size_pt)
    if bold:
        run.bold = True


def _plain_run(paragraph: Any, text: str, size_pt: float | None, bold: bool) -> None:
    if text:
        _style_run(paragraph.add_run(text), size_pt, bold)


def _spacing(paragraph: Any, line: int | None = None, before: int | None = None, after: int | None = None) -> None:
    properties = paragraph._p.get_or_add_pPr()
    spacing = properties.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        properties.append(spacing)
    if line is not None:
        spacing.set(qn("w:line"), str(line))
        spacing.set(qn("w:lineRule"), "auto")
    if before is not None:
        spacing.set(qn("w:before"), str(before))
    if after is not None:
        spacing.set(qn("w:after"), str(after))


def title(document: Any, text: str) -> None:
    """Centered bold 16 pt title.

    Args:
        document: Target document.
        text: Title text.
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _spacing(paragraph, before=240, after=240)
    add_runs(paragraph, text, size_pt=16, bold=True)


def centered(document: Any, text: str, size_pt: float | None = None) -> None:
    """Centered line (authors, affiliations).

    Args:
        document: Target document.
        text: Text with markup.
        size_pt: Optional size.
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(paragraph, text, size_pt=size_pt)


def heading(document: Any, text: str, level: int, numbered: bool) -> None:
    """Heading with the reference numbering (Heading 1 / Heading 2, automatic numbers when `numbered`).

    Args:
        document: Target document.
        text: Heading text.
        level: 1 or 2.
        numbered: Use the automatic multilevel numbering of the reference manuscript.
    """
    paragraph = document.add_paragraph(style=f"Heading {level}")
    _spacing(paragraph, before=0)
    if numbered:
        properties = paragraph._p.get_or_add_pPr()
        numbering = OxmlElement("w:numPr")
        ilvl, num_id = OxmlElement("w:ilvl"), OxmlElement("w:numId")
        ilvl.set(qn("w:val"), str(level - 1))
        num_id.set(qn("w:val"), HEADING_NUM_ID)
        numbering.extend([ilvl, num_id])
        properties.append(numbering)
        indent = OxmlElement("w:ind")
        indent.set(qn("w:left"), "720" if level == 2 else "0")
        indent.set(qn("w:hanging"), "720")
        properties.append(indent)
    run = paragraph.add_run(text)
    run.font.color.rgb = None
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "000000")
    run._r.get_or_add_rPr().append(color)


def body(document: Any, text: str, indent: bool = True) -> None:
    """Double-spaced, justified body paragraph with a first-line indent.

    Args:
        document: Target document.
        text: Text with markup.
        indent: Apply the first-line indent.
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _spacing(paragraph, line=DOUBLE_SPACE)
    if indent:
        paragraph.paragraph_format.first_line_indent = Inches(0.5)
    add_runs(paragraph, text)


def keywords(document: Any, text: str) -> None:
    """Keywords line: bold label, single spaced, justified.

    Args:
        document: Target document.
        text: Keywords separated by semicolons.
    """
    document.add_paragraph()
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(paragraph, "Keywords: ", bold=True)
    add_runs(paragraph, text)


def reference(document: Any, text: str) -> None:
    """Reference entry: double spaced with a 0.5 in hanging indent.

    Args:
        document: Target document.
        text: Reference text with markup (journal names in {i:...}).
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _spacing(paragraph, line=DOUBLE_SPACE)
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    add_runs(paragraph, text)


def caption(document: Any, text: str) -> None:
    """Caption in the reference 'Caption' style (9 pt, justified).

    Args:
        document: Target document.
        text: Caption text with markup.
    """
    paragraph = document.add_paragraph(style="Caption")
    add_runs(paragraph, text)


def figure(document: Any, image: Path, width_in: float) -> None:
    """Centered picture paragraph (the caption follows below).

    Args:
        document: Target document.
        image: PNG path.
        width_in: Picture width in inches.
    """
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image), width=Inches(width_in))


def equation(document: Any, expression: str, number: int) -> None:
    """Centered equation with a right-aligned number, using tab stops.

    Args:
        document: Target document.
        expression: Equation text with markup.
        number: Equation number.
    """
    paragraph = document.add_paragraph()
    _spacing(paragraph, line=240, before=120, after=120)
    stops = paragraph.paragraph_format.tab_stops
    stops.add_tab_stop(Inches(3.25), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    add_runs(paragraph, "\t" + expression + f"\t({number})")


def _cell_borders(cell: Any, top: bool, bottom: bool) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for side, on in (("top", top), ("left", False), ("bottom", bottom), ("right", False)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:val"), "single" if on else "nil")
        if on:
            element.set(qn("w:sz"), "8")
            element.set(qn("w:space"), "0")
            element.set(qn("w:color"), "000000")
        borders.append(element)
    properties.append(borders)


def _shade(cell: Any, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def _margins(cell: Any) -> None:
    margins = OxmlElement("w:tcMar")
    for side, width in (("top", 60), ("left", 80), ("bottom", 60), ("right", 80)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    cell._tc.get_or_add_tcPr().append(margins)


def table(document: Any, header: list[str], rows: list[list[str]], widths_in: list[float], center_from: int = 1) -> None:
    """Table with the reference look: top and bottom rules, grey header, 9 pt text.

    Args:
        document: Target document.
        header: Column titles.
        rows: Row cells (text with markup).
        widths_in: Column widths in inches.
        center_from: Columns from this index on are centered; earlier ones are left aligned.
    """
    grid = document.add_table(rows=1 + len(rows), cols=len(header))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for r, cells in enumerate([header, *rows]):
        for c, text in enumerate(cells):
            cell = grid.cell(r, c)
            cell.width = Inches(widths_in[c])
            _cell_borders(cell, top=r == 0, bottom=r in (0, len(rows)))
            _margins(cell)
            if r == 0:
                _shade(cell, "E8E8E8")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if (c >= center_from and r == 0) or (c >= center_from) else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_runs(paragraph, text, size_pt=9, bold=r == 0)
    spacer = document.add_paragraph()
    _spacing(spacer, after=120)
